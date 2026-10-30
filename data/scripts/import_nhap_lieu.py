"""Chuẩn hóa chuan-bi/mau-nap-du-lieu.xlsx -> CSV nạp bảng tạm import_plans / import_details (D-011).

Chạy (từ thư mục gốc dự án):
  python -I data/scripts/import_nhap_lieu.py --test       # tự kiểm tra các hàm lõi
  python -I data/scripts/import_nhap_lieu.py --mapping    # tạo data/clean/doi-chieu.xlsx (người dùng duyệt)
  python -I data/scripts/import_nhap_lieu.py --build      # đọc doi-chieu.xlsx đã duyệt -> CSV trong data/clean/
"""
import argparse
import bisect
import collections
import csv
import datetime as dt
import difflib
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "chuan-bi" / "mau-nap-du-lieu.xlsx"
CLEAN = ROOT / "data" / "clean"
MAP_FILE = CLEAN / "doi-chieu.xlsx"
DB_FILE = Path(__file__).with_name("danh-muc-hien-co.txt")

# Cột (0-based) của hai sheet nguồn
P_TK, P_DATE, P_ORDER, P_VC, P_CAR, P_DRV, P_TP, P_VT, P_NOTE = range(9)
(B_DATE, B_TK, B_SLIP, B_CT, B_ITEM, B_GIVE, B_RECV, B_LHX, B_MH, B_DVT,
 B_QTY, B_WT, B_NV, B_CARRIER, B_NOTE) = range(15)

SHEETS = ["Phân loại VC", "Loại hình xuất", "Mặt hàng", "ĐVT", "Nhà xe", "NV xuất", "Đơn vị nhận hàng"]
MAP_HEAD = ["Giá trị trong file", "Số dòng", "Đề xuất mục chuẩn", "Có sẵn trên DB?", "Ghi chú"]
GUESS = "đoán — kiểm lại"
CONFIRM = "cần xác nhận"


# ---------------------------------------------------------------- văn bản
def txt(v):
    """Ô -> chuỗi NFC, gộp khoảng trắng; trống / 'Chọn' / chữ tiêu đề cột gõ nhầm vào ô -> None."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(v))).strip()
    return None if s.lower() in ("", "chọn", "loại hình xuất", "mặt hàng") else s


def plain(s, squash=True):
    """Bỏ dấu, hạ chữ, bỏ khoảng trắng, gộp chữ lặp (VÂT TƯ == vật tư)."""
    s = unicodedata.normalize("NFD", s.lower().replace("đ", "d"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"\s+", "", s)
    return re.sub(r"(.)\1+", r"\1", s) if squash else s


def key(s):
    """Khóa so khớp DB: hạ chữ, bỏ khoảng trắng (giữ dấu)."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", s).lower())


def title_vn(s):
    return s.title() if s and (s.isupper() or s.islower()) else s


# ---------------------------------------------------------------- số TK
TK_FIXES = {"TK2004/05-345": "TK2024/05-345"}  # người dùng xác nhận 29/10/2026: gõ sai năm


def norm_tk(raw):
    """Chuẩn hóa rồi áp TK_FIXES (khóa là dạng đã chuẩn hóa)."""
    r = _norm_tk(raw)
    return _norm_tk(TK_FIXES[r[0]]) if r[0] in TK_FIXES else r


def _norm_tk(raw):
    """-> (tk_number, year, month, seq, kind) hoặc (None, ..., 'empty'|'odd')."""
    s = txt(raw)
    if not s:
        return (None, None, None, None, "empty")
    s = s.upper().replace(" ", "")
    m = re.fullmatch(r"TK(\d{4})/(\d{1,2})-(\d{1,3})", s)
    if m:
        y, mo, sq = int(m[1]), int(m[2]), int(m[3])
        if 2000 <= y <= 2100 and 1 <= mo <= 12 and sq >= 1:
            return (f"TK{y}/{mo:02d}-{sq:03d}", y, mo, sq, "new")
    m = re.fullmatch(r"TK(\d{2})/(\d{2})-(\d{4})", s)
    if m:
        y, mo, sq = 2000 + int(m[1]), int(m[2]), int(m[3])
        if 1 <= mo <= 12 and sq >= 1:
            return (s, y, mo, sq, "old")
    return (None, None, None, None, "odd")


# ---------------------------------------------------------------- ngày
def parse_date(v):
    """datetime / 'd/m/yyyy' -> date; None nếu trống; False nếu không đọc được."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*", str(v))
    if m:
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            pass
    return False


def fix_swap(d, tk_ym):
    """Đảo ngày<->tháng CHỈ khi ngày<=12, tháng hiện tại không khớp TK, đảo xong khớp cả tháng và năm của TK."""
    if not d or not tk_ym or d.day > 12 or d.month == tk_ym[1] or d.day != tk_ym[1]:
        return None
    n = dt.date(d.year, d.day, d.month)
    return n if n.year == tk_ym[0] else None


def neighbor_fixes(dates, trusted, cand):
    """Dòng không có TK (cand) mà ngày ≤ 12: đảo ngày/tháng nếu đảo xong nằm giữa hai dòng tin cậy gần nhất
    (file sắp theo ngày) còn giữ nguyên thì không. Không đủ hai bên / không khớp / khớp cả hai -> giữ nguyên.
    -> {chỉ số: ngày mới}"""
    pos = [i for i, t in enumerate(trusted) if t and dates[i]]
    out = {}
    for i, c in enumerate(cand):
        k = bisect.bisect_left(pos, i)
        if not c or k == 0 or k >= len(pos):
            continue
        lo, hi, d = dates[pos[k - 1]], dates[pos[k]], dates[i]
        sw = dt.date(d.year, d.day, d.month)
        if lo <= sw <= hi and not lo <= d <= hi:
            out[i] = sw
    return out


# ---------------------------------------------------------------- số
def parse_number(v):
    """-> (giá trị float|None, quy tắc|None). quy tắc None = ô số gốc, không cần ghi log."""
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        return (None, "số âm -> trống") if v < 0 else (float(v), None)
    s = re.sub(r"\s+", "", str(v))
    if s == "":
        return None, None
    if s.startswith("-"):
        return None, "số âm -> trống"
    if re.fullmatch(r"\d+", s):
        return float(s), "số dạng chữ"
    if re.fullmatch(r"[1-9]\d{0,2}(\.\d{3})+(,\d+)?", s):
        amb = "," not in s
        return float(s.replace(".", "").replace(",", ".")), "VN" + (" (1 dấu chấm + 3 số: nghìn, không chắc)" if amb and s.count(".") == 1 else "")
    if re.fullmatch(r"\d+,\d+", s):
        return float(s.replace(",", ".")), "VN thập phân" + (" (3 số sau phẩy, không chắc)" if len(s.split(",")[1]) == 3 else "")
    if re.fullmatch(r"\d{1,3}(,\d{3})+\.\d+", s) or re.fullmatch(r"\d{1,3}(,\d{3}){2,}", s):
        return float(s.replace(",", "")), "kiểu Mỹ (phẩy nghìn, chấm thập phân)"
    if re.fullmatch(r"\d+\.\d+", s):
        return float(s), "chấm thập phân (không theo nhóm 3 số)"
    return None, "không đọc được -> trống"


def fmt_num(x):
    if x is None:
        return ""
    s = f"{round(x, 3):.3f}".rstrip("0").rstrip(".")
    return "0" if s in ("", "-0") else s


# ---------------------------------------------------------------- đọc nguồn & DB
def read_sheet(wb, name, ncol):
    """-> [(số dòng excel, [ô...])] từ dòng 3, bỏ dòng hoàn toàn trống."""
    out, blank = [], 0
    for i, r in enumerate(wb[name].iter_rows(values_only=True), 1):
        if i < 3:
            continue
        c = (list(r) + [None] * ncol)[:ncol]
        if all(txt(x) is None for x in c):
            blank += 1
        else:
            out.append((i, c))
    return out, blank


def load_source():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    P, pb = read_sheet(wb, "Nhap phieu", 9)
    B, bb = read_sheet(wb, "Bang kiem", 15)
    return P, B, pb, bb


def load_db():
    kinds = {"transport_type", "export_type", "product", "unit", "carrier", "delivery_unit", "fg_exporter"}
    db = {}
    for line in DB_FILE.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^(\w+)[^:]*:\s*(.*)$", line)
        if m and m[1] in kinds:
            db[m[1]] = [x.strip() for x in m[2].split(" | ") if x.strip()]
    return db


# ---------------------------------------------------------------- đề xuất đối chiếu
# (giá trị đã bỏ dấu/khoảng trắng/chữ lặp -> (đề xuất|None, ghi chú))
S_PRODUCT = {
    plain("vật tư"): ("9 - Vật tư phụ, BHLĐ, VLH", "gộp các biến thể 'vật tư'"),
    plain("vt"): ("9 - Vật tư phụ, BHLĐ, VLH", GUESS + ": VT = vật tư"),
    plain("vlh"): ("9 - Vật tư phụ, BHLĐ, VLH", GUESS + ": VLH nằm trong tên mục 9"),
    plain("bhlđ"): ("9 - Vật tư phụ, BHLĐ, VLH", GUESS + ": mã 10 cũ = BHLĐ, trên DB mã 10 là Sơn"),
    plain("kct"): ("21 - KCT TD", GUESS),
    plain("tole"): ("18 - TOLE TD", GUESS),
    plain("bulong/vít"): ("12 - Bulong/Vít /đinh hàn", GUESS + ": mã 16 cũ, trên DB mã 16 là Loại khác"),
    plain("nhớt"): ("11 - Nhiên liệu (dầu, nhớt)", GUESS),
    plain("nhiên liệu"): ("11 - Nhiên liệu (dầu, nhớt)", GUESS),
    plain("điều chuyển nội bộ"): (None, "giá trị của cột Loại hình xuất nằm nhầm cột Mặt hàng — để trống; " + GUESS),
}
S_EXPORT = {
    plain("xuất thuê/mượn"): ("12 - Xuất thuê-mượn", GUESS + ": mã 8 trong file, trên DB mã 8 là Xuất đổi hàng"),
}
S_VC = {
    plain("bán"): ("An Hạ -> Xuất Bán", GUESS),
    plain("gcn"): ("An Hạ -> GCN", GUESS),
}
S_CARRIER = {
    plain("lê thuận"): (None, "tên người, không phải nhà xe — để trống; " + CONFIRM),
    plain("0"): (None, "giá trị 0 — để trống"),
    plain("nhà xe"): (None, "chữ tiêu đề lẫn trong dữ liệu — để trống"),
    plain("xe thuê"): ("Xe thuê", "giữ nguyên; có thể cùng nghĩa 'Xe ngoài DDC' — " + CONFIRM),
}
UNIT_FIX = {  # plain -> (chuẩn, ghi chú)
    "pcs": ("PCS", ""), "psc": ("PCS", "PSC gõ sai PCS"), "kg": ("Kg", ""), "m": ("M", ""),
    "met": ("M", GUESS + ": Mét = M"), "m2": ("M2", ""), "m3": ("M3", ""), "mm": ("mm", ""),
    "lit": ("Lít", ""), "liit": ("Lít", GUESS + ": LIÍT gõ sai Lít"), "set": ("Set", ""),
    "soi": ("Sợi", GUESS + ": Sơi gõ sai Sợi"), "tuyt": ("Tuýp", GUESS + ": tuýt gõ sai Tuýp"),
}


def catalog_proposal(v, db_names, special, numbered=False, vc=False):
    """-> (đề xuất|None, ghi chú, đã-khớp-DB?)."""
    s = re.sub(r"\s*>>\s*", " -> ", v) if vc else v
    m = re.fullmatch(r"(\d+)\s*-\s*(.+)", s) if numbered else None
    norm = f"{m[1]} - {m[2].strip()}" if m else s
    byfull = {key(n): n for n in db_names}
    if key(norm) in byfull:
        return byfull[key(norm)], "", True
    name = m[2].strip() if m else s
    if numbered:
        byname = {key(re.sub(r"^\d+\s*-\s*", "", n)): n for n in db_names}
        if key(name) in byname:
            hit = byname[key(name)]
            return hit, ("" if not m else f"mã số trong file ({m[1]}) khác DB — kiểm lại"), True
    if plain(name) in special:
        p, note = special[plain(name)]
        return p, note, p in db_names
    return norm, "", False


def unit_proposal(v):
    if re.fullmatch(r"[\d.,]+", v):
        return None, "ô số nằm nhầm cột ĐVT — để trống; " + CONFIRM
    k = plain(v, squash=False)
    if k == "giacongngoai":
        return None, "chữ của cột Loại hình xuất nằm nhầm cột ĐVT — để trống; " + CONFIRM
    if k in UNIT_FIX:
        return UNIT_FIX[k]
    return v.lower().capitalize(), ""


def person_proposal(v, db_names):
    """Tên người: khớp DB không phân biệt hoa/thường/dấu; gần đúng thì ghi 'đoán'."""
    byplain = {plain(n): n for n in db_names}
    if plain(v) in byplain:
        hit = byplain[plain(v)]
        return hit, ("" if hit.lower() == v.lower() else "khác dấu so với DB — kiểm lại"), True
    pv = unicodedata.normalize("NFD", v.lower().replace("đ", "d"))
    pv = "".join(c for c in pv if unicodedata.category(c) != "Mn").split()
    for n in db_names:
        pn = plain(n)
        if pv and plain(pv[-1]) == plain(n.split()[-1]) and difflib.SequenceMatcher(None, plain(v), pn).ratio() >= 0.8:
            return n, GUESS + ": viết tắt/sai chính tả của tên trên DB", True
    return title_vn(v), "không có trên DB — sẽ tạo mới; " + CONFIRM, False


def merge_case(rows):
    """Các mục chưa có trên DB mà chỉ khác hoa/thường -> dùng dạng xuất hiện nhiều nhất."""
    groups = collections.defaultdict(list)
    for r in rows:
        if r["prop"] and r["indb"] is not True:
            groups[key(r["prop"])].append(r)
    for g in groups.values():
        best = max(g, key=lambda r: r["count"])["prop"]
        for r in g:
            if r["prop"] != best:
                r["prop"] = best
                r["note"] = (r["note"] + "; " if r["note"] else "") + "gộp biến thể hoa/thường"


def build_rows(counter, fn):
    rows = []
    for v, n in sorted(counter.items(), key=lambda x: (-x[1], x[0])):
        prop, note, indb = fn(v)
        rows.append({"v": v, "count": n, "prop": prop, "note": note, "indb": indb})
    merge_case(rows)
    return rows


def split_names(raw):
    return [p for p in (txt(x) for x in str(raw).split(";")) if p] if raw is not None else []


def collect(P, B):
    """Đếm giá trị từng danh mục cần duyệt."""
    cnt = {s: collections.Counter() for s in SHEETS}
    for _, c in P:
        if txt(c[P_VC]):
            cnt["Phân loại VC"][txt(c[P_VC])] += 1
        for col in (P_TP, P_VT):
            for part in split_names(c[col]):
                cnt["NV xuất"][part] += 1
    for _, c in B:
        for col, s in ((B_LHX, "Loại hình xuất"), (B_MH, "Mặt hàng"), (B_DVT, "ĐVT"), (B_CARRIER, "Nhà xe"),
                       (B_RECV, "Đơn vị nhận hàng")):
            if txt(c[col]):
                cnt[s][txt(c[col])] += 1
        for part in split_names(c[B_NV]):
            cnt["NV xuất"][part] += 1
    return cnt


def make_mapping(force):
    if MAP_FILE.exists() and not force:
        sys.exit(f"{MAP_FILE} đã tồn tại (có thể đã được duyệt). Chạy lại với --force nếu muốn ghi đè.")
    P, B, _, _ = load_source()
    db = load_db()
    cnt = collect(P, B)
    fns = {
        "Phân loại VC": lambda v: catalog_proposal(v, db["transport_type"], S_VC, vc=True),
        "Loại hình xuất": lambda v: catalog_proposal(v, db["export_type"], S_EXPORT, numbered=True),
        "Mặt hàng": lambda v: catalog_proposal(v, db["product"], S_PRODUCT, numbered=True),
        "ĐVT": lambda v: (lambda p, n: (p, n, p in db["unit"] if p else False))(*unit_proposal(v)),
        "Nhà xe": lambda v: catalog_proposal(v, db["carrier"], S_CARRIER),
        "NV xuất": lambda v: person_proposal(v, db["fg_exporter"]),
        "Đơn vị nhận hàng": lambda v: (v, "", None),
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hướng dẫn"
    for line in [
        "BẢNG ĐỐI CHIẾU DANH MỤC — duyệt trước khi nạp dữ liệu cũ",
        "Mỗi sheet là một danh mục. Chỉ sửa cột C 'Đề xuất mục chuẩn' (cột A, B, D, E chỉ để tham khảo).",
        "Cột C để TRỐNG = trường đó để trống khi nạp. Hai dòng cùng ghi một mục ở cột C = gộp làm một.",
        "Ô màu vàng = đề xuất dựa trên đoán hoặc cần bạn xác nhận. Cột D: 'Có' = mục đã có trên hệ thống, 'Không' = sẽ được tạo mới.",
        "Số xe, tài xế, mã công trình không đưa vào đây (chỉ chuẩn hóa khoảng trắng; tài xế viết hoa -> Title Case).",
        "Lưu file (giữ nguyên tên và tên sheet) rồi báo lại để chạy bước --build.",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for s in SHEETS:
        rows = build_rows(cnt[s], fns[s])
        w = wb.create_sheet(s)
        w.append(MAP_HEAD)
        for r in rows:
            indb = "không rõ" if r["indb"] is None else ("Có" if r["indb"] else ("Không" if r["prop"] else ""))
            note = r["note"]
            if r["prop"] and r["indb"] is False and CONFIRM not in note and "sẽ tạo mới" not in note:
                note = (note + "; " if note else "") + "chưa có trên DB — sẽ tạo mới" + ("" if s == "ĐVT" else "; " + CONFIRM)
            w.append([r["v"], r["count"], r["prop"], indb, note])
            for cell in w[w.max_row]:
                if cell.data_type == "f":  # giá trị bắt đầu bằng '=' không được thành công thức
                    cell.data_type = "s"
            if GUESS in note or CONFIRM in note:
                w.cell(w.max_row, 3).fill = yellow
        for col, width in zip("ABCDE", (42, 10, 42, 16, 70)):
            w.column_dimensions[col].width = width
        for cell in w[1]:
            cell.font = Font(bold=True)
        w.freeze_panes = "A2"
        w.auto_filter.ref = f"A1:E{w.max_row}"
        print(f"  {s}: {len(rows)} giá trị, {sum(r['count'] for r in rows)} lượt dùng, "
              f"{sum(1 for r in rows if GUESS in r['note'] or CONFIRM in r['note'])} dòng tô vàng")
    CLEAN.mkdir(parents=True, exist_ok=True)
    wb.save(MAP_FILE)
    print(f"Đã ghi {MAP_FILE}")


class Maps:
    def __init__(self):
        wb = openpyxl.load_workbook(MAP_FILE, read_only=True, data_only=True)
        self.d, self.missing, self.missing_vals = {}, collections.Counter(), collections.Counter()
        for s in SHEETS:
            rows = list(wb[s].iter_rows(values_only=True))
            if [txt(x) for x in rows[0][:5]] != MAP_HEAD:
                sys.exit(f"Sheet '{s}' trong {MAP_FILE.name} sai tiêu đề cột")
            self.d[s] = {txt(r[0]): txt(r[2]) for r in rows[1:] if txt(r[0]) is not None}

    def get(self, sheet, raw):
        v = txt(raw)
        if v is None:
            return None
        if v in self.d[sheet]:
            return self.d[sheet][v]
        self.missing[sheet] += 1
        self.missing_vals[(sheet, v)] += 1
        return v

    def names(self, raw):
        out = []
        for part in split_names(raw):
            n = self.get("NV xuất", part)
            if n and n not in out:
                out.append(n)
        return ";".join(out) or None


# ---------------------------------------------------------------- build
def write_csv(path, header, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:  # UTF-8 không BOM
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def build():
    selftest()
    if not MAP_FILE.exists():
        sys.exit(f"Chưa có {MAP_FILE}; chạy --mapping trước.")
    P, B, pblank, bblank = load_source()
    maps = Maps()
    st = collections.Counter()
    rejected, date_fixes, num_fixes, odd_tk = [], [], [], []
    st["P dòng trống bỏ qua"], st["B dòng trống bỏ qua"] = pblank, bblank

    def reject(sheet, row, c, tk, reason):
        rejected.append([sheet, row, tk or "", reason, " | ".join(txt(x) or "" for x in c)])
        st[f"loại {sheet}: {reason.split(':')[0]}"] += 1

    def get_date(sheet, row, v, tkinfo):
        d = parse_date(v)
        if d is False:
            date_fixes.append([sheet, row, tkinfo[0] or "", str(v), "", "không đọc được -> trống", ""])
            st[f"{sheet}: ngày không đọc được"] += 1
            return None
        if d is None:
            st[f"{sheet}: ngày trống"] += 1
            return None
        n = fix_swap(d, (tkinfo[1], tkinfo[2]) if tkinfo[0] else None)
        if n:
            date_fixes.append([sheet, row, tkinfo[0], d.isoformat(), n.isoformat(), "đảo ngày/tháng", "theo số TK"])
            st[f"{sheet}: ngày đã đảo theo số TK"] += 1
            return n
        return d

    # --- phiếu
    plans, seen = [], {}
    for row, c in P:
        tk = norm_tk(c[P_TK])
        if tk[4] == "empty":
            reject("Nhap phieu", row, c, None, "thiếu số TK")
            continue
        if tk[4] == "odd":
            reject("Nhap phieu", row, c, txt(c[P_TK]), f"số TK dạng lạ: {txt(c[P_TK])}")
            continue
        if tk[0] in seen:
            reject("Nhap phieu", row, c, tk[0], f"trùng số TK: giống dòng {seen[tk[0]]}")
            continue
        seen[tk[0]] = row
        d = get_date("Nhap phieu", row, c[P_DATE], tk)
        plans.append({
            "tk_number": tk[0], "tk_year": tk[1], "tk_month": tk[2], "tk_seq": tk[3],
            "export_date": d, "transport_order_no": txt(c[P_ORDER]),
            "transport_type": maps.get("Phân loại VC", c[P_VC]), "vehicle": txt(c[P_CAR]),
            "driver": title_vn(txt(c[P_DRV])), "fg_exporter": maps.names(c[P_TP]),
            "material_exporter": maps.names(c[P_VT]), "note": txt(c[P_NOTE]), "is_stub": False, "source_row": row,
        })
    st["phiếu nạp (từ Nhap phieu)"] = len(plans)

    # --- bảng kiểm
    details, stub_dates, stub_info = [], collections.defaultdict(list), {}
    trusted, cand = [], []  # cho neighbor_fixes
    for row, c in B:
        for col in (B_LHX, B_MH):
            if isinstance(c[col], str) and txt(c[col]) is None and txt(c[col] + "x") != "chọnx":
                st["B ô chữ tiêu đề cột gõ nhầm (coi như trống)"] += 1
        tk = norm_tk(c[B_TK])
        note = txt(c[B_NOTE])
        if tk[4] == "odd":
            odd_tk.append(["Bang kiem", row, txt(c[B_TK]), "coi như không có TK; TK gốc ghi vào note"])
            note = "; ".join(x for x in (note, f"TK gốc dạng lạ: {txt(c[B_TK])}") if x)
            st["B dòng có TK dạng lạ (coi như không có TK)"] += 1
        elif tk[0] and _norm_tk(c[B_TK])[0] in TK_FIXES:
            odd_tk.append(["Bang kiem", row, txt(c[B_TK]), f"sửa thành {tk[0]} (người dùng xác nhận)"])
            st["B dòng đổi số TK theo TK_FIXES"] += 1
        d0 = parse_date(c[B_DATE])
        d = get_date("Bang kiem", row, c[B_DATE], tk)
        amb = isinstance(c[B_DATE], dt.datetime) and bool(d0) and d0.day <= 12 and d0.day != d0.month
        trusted.append(bool(d) and (d != d0 or not amb or bool(tk[0] and d.month == tk[2])))
        cand.append(bool(d) and amb and not tk[0])
        nums = []
        for col, name in ((B_QTY, "Số lượng xuất"), (B_WT, "Tổng khối lượng (kg)")):
            val, rule = parse_number(c[col])
            if rule:
                num_fixes.append(["Bang kiem", row, name, str(c[col]), fmt_num(val), rule])
                st[f"số: {rule.split(' (')[0]}"] += 1
            nums.append(fmt_num(val))
        if tk[0]:
            st["B dòng có TK"] += 1
            if tk[0] not in seen:
                stub_info.setdefault(tk[0], tk)
                if d:
                    stub_dates[tk[0]].append(d)
        else:
            st["B dòng không có TK"] += 1
        details.append([
            f"BK-{row}", tk[0], d, txt(c[B_SLIP]), txt(c[B_CT]), txt(c[B_ITEM]),
            txt(c[B_GIVE]) or "An Hạ", maps.get("Đơn vị nhận hàng", c[B_RECV]),
            maps.get("Loại hình xuất", c[B_LHX]), maps.get("Mặt hàng", c[B_MH]), maps.get("ĐVT", c[B_DVT]),
            nums[0], nums[1], maps.names(c[B_NV]), maps.get("Nhà xe", c[B_CARRIER]), note, row,
        ])
    st["dòng bảng kiểm nạp"] = len(details)

    # --- đảo ngày theo dòng lân cận (dòng không có TK)
    new = neighbor_fixes([r[2] for r in details], trusted, cand)
    for i, n in new.items():
        date_fixes.append(["Bang kiem", details[i][-1], "", details[i][2].isoformat(), n.isoformat(),
                           "đảo ngày/tháng", "theo dòng lân cận"])
        details[i][2] = n
    st["Bang kiem: ngày đã đảo theo dòng lân cận"] = len(new)
    st["Bang kiem: ứng viên (không TK, ngày ≤ 12) giữ nguyên"] = sum(cand) - len(new)
    today = dt.date(2026, 10, 29)
    st["Bang kiem: dòng ngày > 2026-10-29"] = sum(1 for r in details if r[2] and r[2] > today)
    st["Nhap phieu: phiếu ngày > 2026-10-29"] = sum(1 for p in plans if p["export_date"] and p["export_date"] > today)

    # --- phiếu tự tạo
    stubs = []
    for t, tk in sorted(stub_info.items(), key=lambda x: (x[1][1], x[1][2], x[1][3])):
        ds = stub_dates.get(t)
        stubs.append({"tk_number": t, "tk_year": tk[1], "tk_month": tk[2], "tk_seq": tk[3],
                      "export_date": min(ds) if ds else None, "transport_order_no": None, "transport_type": None,
                      "vehicle": None, "driver": None, "fg_exporter": None, "material_exporter": None,
                      "note": None, "is_stub": True, "source_row": ""})
    st["phiếu tự tạo"] = len(stubs)
    st["  trong đó số TK kiểu cũ"] = sum(1 for s in stubs if re.match(r"TK\d{2}/", s["tk_number"]))

    # --- va chạm (năm, tháng, seq)
    by_ym = collections.defaultdict(list)
    for p in plans + stubs:
        by_ym[(p["tk_year"], p["tk_month"], p["tk_seq"])].append(p)
    collisions = [[*k, " ; ".join(f"{p['tk_number']}{' (tự tạo)' if p['is_stub'] else ''}" for p in v)]
                  for k, v in sorted(by_ym.items()) if len(v) > 1]
    st["va chạm (năm,tháng,seq)"] = len(collisions)

    # --- ghi file
    CLEAN.mkdir(parents=True, exist_ok=True)
    pcols = ["tk_number", "tk_year", "tk_month", "tk_seq", "export_date", "transport_order_no", "transport_type",
             "vehicle", "driver", "fg_exporter", "material_exporter", "note", "is_stub", "source_row"]

    def cell(x):
        return "" if x is None else ("true" if x is True else "false" if x is False else
                                     x.isoformat() if isinstance(x, dt.date) else str(x))

    write_csv(CLEAN / "import_plans.csv", pcols, [[cell(p[k]) for k in pcols] for p in plans + stubs])
    write_csv(CLEAN / "import_details.csv",
              ["import_ref", "tk_number", "export_date", "warehouse_slip_no", "construction_code", "work_item",
               "delivery_unit", "receiving_unit", "export_type", "product", "unit", "quantity", "total_weight_kg",
               "shipper", "carrier", "note", "source_row"], [[cell(x) for x in r] for r in details])
    write_csv(CLEAN / "_rejected.csv", ["sheet", "source_row", "tk_number", "reason", "raw_values"], rejected)
    write_csv(CLEAN / "_date_fixes.csv", ["sheet", "source_row", "tk_number", "original", "fixed", "kind", "reason"], date_fixes)
    write_csv(CLEAN / "_number_fixes.csv", ["sheet", "source_row", "column", "original", "converted", "rule"], num_fixes)
    write_csv(CLEAN / "_tk_collisions.csv", ["tk_year", "tk_month", "tk_seq", "tk_numbers"], collisions)
    write_csv(CLEAN / "_tk_odd.csv", ["sheet", "source_row", "tk_original", "action"], odd_tk)

    # --- đối chiếu số dòng
    assert len(plans) + sum(1 for r in rejected if r[0] == "Nhap phieu") == len(P), "Nhap phieu: nạp + loại != đọc"
    assert len(details) + sum(1 for r in rejected if r[0] == "Bang kiem") == len(B), "Bang kiem: nạp + loại != đọc"
    assert len({p["tk_number"] for p in plans + stubs}) == len(plans) + len(stubs), "tk_number trùng"
    assert len({d[0] for d in details}) == len(details), "import_ref trùng"
    print(f"Nhap phieu: đọc {len(P)} (+{pblank} dòng trống) = nạp {len(plans)} + loại "
          f"{sum(1 for r in rejected if r[0] == 'Nhap phieu')}")
    print(f"Bang kiem : đọc {len(B)} (+{bblank} dòng trống) = nạp {len(details)} + loại "
          f"{sum(1 for r in rejected if r[0] == 'Bang kiem')}")
    for k, v in sorted(st.items()):
        print(f"  {k}: {v}")
    if maps.missing:
        print("CẢNH BÁO giá trị không có trong doi-chieu.xlsx (giữ nguyên):", dict(maps.missing))
        write_csv(CLEAN / "_new_catalog_values.csv", ["sheet", "value", "rows"],
                  [[s, v, n] for (s, v), n in sorted(maps.missing_vals.items())])
    tot = sum(float(d[12]) for d in details if d[12])
    print(f"  tổng khối lượng (kg) trong import_details.csv: {tot:,.3f}")


# ---------------------------------------------------------------- test
def selftest():
    assert norm_tk("TK2024/3-324")[:4] == ("TK2024/03-324", 2024, 3, 324)
    assert norm_tk(" tk2023/12-5 ")[0] == "TK2023/12-005"
    assert norm_tk("tk23/01-0001")[:4] == ("TK23/01-0001", 2023, 1, 1)
    assert norm_tk("TK23/01-0001A")[4] == "odd" and norm_tk("TK2023/13-001")[4] == "odd"
    assert norm_tk(None)[4] == "empty"
    d = dt.date
    assert fix_swap(d(2023, 1, 5), (2023, 5)) == d(2023, 5, 1)
    assert fix_swap(d(2023, 5, 1), (2023, 5)) is None  # đã khớp
    assert fix_swap(d(2023, 1, 15), (2023, 5)) is None  # ngày > 12
    assert fix_swap(d(2024, 1, 12), (2023, 12)) is None  # đảo xong lệch năm TK
    assert fix_swap(d(2023, 1, 5), None) is None
    assert norm_tk("TK2004/5-345")[:4] == ("TK2024/05-345", 2024, 5, 345)
    nf = lambda ds, t, c: neighbor_fixes(ds, t, c)
    ds = [d(2023, 5, 1), d(2023, 1, 5), d(2023, 5, 3), d(2023, 1, 6)]
    assert nf(ds, [1, 0, 1, 0], [0, 1, 0, 1]) == {1: d(2023, 5, 1)}  # dòng cuối thiếu lân cận phải -> giữ
    ds = [d(2023, 1, 1), d(2023, 1, 5), d(2023, 1, 20)]
    assert nf(ds, [1, 0, 1], [0, 1, 0]) == {}  # giữ nguyên khớp, đảo không khớp
    ds = [d(2023, 5, 1), d(2023, 9, 3), d(2023, 5, 20)]
    assert nf(ds, [1, 0, 1], [0, 1, 0]) == {}  # không khớp cả hai
    assert parse_date("05/01/2023") == d(2023, 1, 5) and parse_date("5/1/2023") == d(2023, 1, 5)
    assert parse_date("#REF!") is False and parse_date(None) is None
    pn = lambda s: parse_number(s)[0]
    assert pn("1.234,5") == 1234.5 and pn("5,5") == 5.5 and pn("1.234") == 1234 and pn("  5,00 ") == 5
    assert pn("44.2") == 44.2 and pn("05.09") == 5.09 and pn("6.00") == 6 and pn("0.500") == 0.5
    assert pn("32,011.00") == 32011 and pn("1,234,567") == 1234567
    assert pn("\\") is None and pn("17,332,8") is None and pn(-3) is None and pn(7) == 7.0
    assert parse_number(7)[1] is None
    assert fmt_num(12.3456) == "12.346" and fmt_num(5.0) == "5" and fmt_num(None) == ""
    assert title_vn("NGUYỄN VĂN SÁU") == "Nguyễn Văn Sáu" and title_vn("Phạm Hữu Cường") == "Phạm Hữu Cường"
    assert txt("  a   b ") == "a b" and txt("Chọn") is None
    print("selftest OK")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--mapping", action="store_true")
    ap.add_argument("--build", action="store_true")
    ap.add_argument("--test", action="store_true")
    ap.add_argument("--force", action="store_true", help="ghi đè doi-chieu.xlsx đã có")
    a = ap.parse_args()
    if a.test:
        selftest()
    elif a.mapping:
        make_mapping(a.force)
    elif a.build:
        build()
    else:
        ap.print_help()


if __name__ == "__main__":
    main()




